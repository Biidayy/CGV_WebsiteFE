#!/usr/bin/env python3
"""Generate complete Word report from research PDF + User Guide."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path("/workspace")
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "docs" / "BAO-CAO-HOAN-CHINH-CGV-WEBSITE.docx"
ARTIFACT = Path("/opt/cursor/artifacts/BAO-CAO-HOAN-CHINH-CGV-WEBSITE.docx")

FONT = "Times New Roman"
RED = RGBColor(0xE7, 0x1A, 0x0F)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_run_font(run, size=12, bold=False, color=DARK, font=FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_para(doc, text, *, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             space_after=6, space_before=0, first_indent=False, color=DARK):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(1)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    sizes = {1: 16, 2: 14, 3: 12}
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 12), bold=True, color=RED if level == 1 else DARK)
    return p


def add_bullets(doc, items, size=12):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        # clear default runs then add styled
        if p.runs:
            p.runs[0].text = item
            set_run_font(p.runs[0], size=size)
        else:
            run = p.add_run(item)
            set_run_font(run, size=size)


def set_cell_text(cell, text, *, bold=False, size=11, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def shade_header_row(row):
    from docx.oxml import OxmlElement
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2D7D5")
        tcPr.append(shd)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=11, center=True)
    shade_header_row(table.rows[0])
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i + 1].cells[c_i], str(val), size=11)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_image(doc, path, caption, width_inches=5.8):
    path = Path(path)
    if not path.exists():
        add_para(doc, f"[Thiếu ảnh: {path.name}]", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    set_run_font(r, size=11, bold=True)


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()

    # Page setup A4
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # ========== COVER ==========
    for _ in range(2):
        doc.add_paragraph()
    add_para(doc, "TRƯỜNG ĐẠI HỌC VĂN LANG", size=14, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "KHOA CÔNG NGHỆ THÔNG TIN", size=13, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    add_para(doc, "BÁO CÁO ĐỒ ÁN", size=18, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, color=RED)
    add_para(doc, "THIẾT KẾ GIAO DIỆN NGƯỜI DÙNG", size=16, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    add_para(doc, "Nghiên cứu và tái hiện giao diện website CGV Cinemas Vietnam",
             size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=28)

    add_para(doc, "Nhóm: 6", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "Giảng viên hướng dẫn: Thầy Nguyễn Hồng Diên", size=12,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    add_table(doc,
              ["Họ và tên", "MSSV", "Vai trò"],
              [
                  ["Hồ Huy", "2374802010778", "Nghiên cứu UI, trang chính, FE, User Guide"],
                  ["Phan Văn Khải", "2474802010164", "Trang mở rộng, CSS/JS, review code"],
              ],
              col_widths=[4.5, 3.5, 7])

    add_para(doc, "Website: https://cgv-website-fe.vercel.app", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "GitHub: https://github.com/Biidayy/CGV_WebsiteFE", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "Domain: https://cgv-website-tkgd.is-a.dev", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    add_para(doc, "TP. Hồ Chí Minh, 2026", size=12, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER)

    page_break(doc)

    # ========== MỤC LỤC ==========
    add_heading_custom(doc, "MỤC LỤC", 1)
    toc = [
        "1. Giới thiệu đồ án",
        "2. Kế hoạch triển khai và phân công",
        "3. Nghiên cứu thói quen người dùng",
        "4. Cấu trúc website",
        "5. Hệ thống liên kết và tương tác",
        "6. Màu sắc và hình thức thể hiện",
        "7. Sitemap và User Flow",
        "8. Wireframe và thiết kế các trang chính",
        "9. Tính năng nổi bật (20+ tính năng tự phát triển)",
        "10. Hướng dẫn sử dụng (User Guide)",
        "11. Ảnh chụp giao diện (Screenshots)",
        "12. Giới hạn của hệ thống",
        "13. Kết luận",
        "Tài liệu tham khảo / Liên kết nộp bài",
    ]
    for item in toc:
        add_para(doc, item, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)

    page_break(doc)

    # ========== 1 ==========
    add_heading_custom(doc, "1. Giới thiệu đồ án", 1)
    add_para(doc,
             "Đồ án thuộc môn Thiết kế Giao diện Người dùng, nhằm nghiên cứu và tái hiện "
             "giao diện website CGV Cinemas Vietnam trên nền tảng Front-End tĩnh. Website "
             "bám sát cấu trúc điều hướng của cgv.vn, đồng thời bổ sung các tính năng tương "
             "tác phía client (tìm kiếm, yêu thích, đặt vé mock, localStorage) để phục vụ "
             "mục tiêu học tập và báo cáo.",
             first_indent=True)
    add_para(doc,
             "Stack công nghệ: HTML5 + CSS3 + JavaScript. Không sử dụng backend, không thanh "
             "toán thật, không đăng nhập thật. Dữ liệu phim/rạp/suất là mock data trong "
             "js/data.js. Website đã được deploy công khai qua Vercel / GitHub Pages.",
             first_indent=True)

    add_heading_custom(doc, "1.1. Mục tiêu", 2)
    add_bullets(doc, [
        "Nghiên cứu UI/UX và thói quen người dùng trên website gốc cgv.vn.",
        "Thiết kế sitemap, user flow, palette màu và wireframe trên Figma.",
        "Hiện thực Front-End tĩnh với các trang chính và luồng đặt vé mock.",
        "Bổ sung 20+ tính năng tương tác, responsive, lưu trữ localStorage.",
        "Deploy công khai và cung cấp User Guide đầy đủ kèm screenshot.",
    ])

    # ========== 2 ==========
    add_heading_custom(doc, "2. Kế hoạch triển khai và phân công", 1)
    add_heading_custom(doc, "2.1. Timeline dự án", 2)
    add_table(doc,
              ["Giai đoạn", "Nội dung công việc"],
              [
                  ["Tuần 1–2", "Nghiên cứu cgv.vn, Figma, thói quen người dùng"],
                  ["Tuần 2–3", "Thiết kế wireframe, sitemap, user flow, màu/font"],
                  ["Tuần 3–6", "Code Front-End HTML/CSS/JS — trang chính & flow đặt vé"],
                  ["Tuần 6–7", "Tích hợp branch, responsive, localStorage"],
                  ["Tuần 7–8", "Deploy GitHub Pages/Vercel, User Guide, báo cáo"],
              ],
              col_widths=[3.5, 12])

    add_heading_custom(doc, "2.2. Phân công thành viên", 2)
    add_table(doc,
              ["Thành viên", "Nhánh Git", "Công việc", "Tiến độ"],
              [
                  ["Hồ Huy", "huy/main", "Nghiên cứu UI, trang chính Figma, tính năng FE, User Guide", "100%"],
                  ["Phan Văn Khải", "Khai-code/main", "Trang mở rộng, CSS/JS, review code", "100%"],
              ],
              col_widths=[3.5, 3.2, 7.3, 2])

    # ========== 3 ==========
    add_heading_custom(doc, "3. Nghiên cứu thói quen người dùng", 1)
    add_para(doc,
             "Qua quan sát website gốc và hành vi đặt vé điện ảnh, nhóm xác định bốn giai "
             "đoạn chính trong hành trình người dùng:",
             first_indent=True)
    add_table(doc,
              ["Giai đoạn", "Hành vi người dùng"],
              [
                  ["Khám phá & Tìm kiếm", "Quét banner khuyến mãi & carousel phim; tìm theo tên/thể loại/định dạng (IMAX, 4DX)"],
                  ["So sánh & Quyết định", "Đối chiếu phim đang chiếu vs sắp chiếu; xem ưu đãi trước khi đặt vé"],
                  ["Chọn rạp & Đặt vé", "Ưu tiên rạp gần nhất theo tỉnh/thành; bấm “Mua vé” nhanh chóng"],
                  ["Sau đặt vé", "Kiểm tra “Vé của tôi”, lưu phim yêu thích; ưu tiên trải nghiệm mobile-first"],
              ],
              col_widths=[4, 11.5])

    # ========== 4 ==========
    add_heading_custom(doc, "4. Cấu trúc website", 1)
    add_para(doc, "Stack: HTML5 + CSS3 + JavaScript (Front-End tĩnh, không backend).", first_indent=True)
    add_table(doc,
              ["Thư mục / File", "Mô tả"],
              [
                  ["cgv/*.html", "Trang chủ, phim, rạp, đặt vé, ưu đãi và các trang mở rộng"],
                  ["cgv/css/", "style.css (layout chung) + CSS theo từng trang"],
                  ["cgv/js/data.js", "Mock data: phim, rạp, suất chiếu, ưu đãi"],
                  ["cgv/js/main.js", "Render Header/Footer dùng chung, UI, filter, booking"],
                  ["cgv/js/storage.js", "Helpers localStorage (favorites, tickets, recent)"],
                  ["docs/", "Screenshot & tài liệu kỹ thuật / báo cáo"],
                  [".github/workflows/", "Auto deploy CI/CD GitHub Pages"],
                  ["vercel.json", "Cấu hình rewrite deploy Vercel"],
              ],
              col_widths=[4.5, 11])
    add_para(doc,
             "Header & Footer dùng chung trên toàn bộ trang, được render động bởi main.js. "
             "Website được triển khai và truy cập qua domain Vercel "
             "(https://cgv-website-fe.vercel.app) và GitHub Pages.",
             first_indent=True)

    # ========== 5 ==========
    add_heading_custom(doc, "5. Hệ thống liên kết và tương tác", 1)
    add_heading_custom(doc, "5.1. Điều hướng chính", 2)
    add_bullets(doc, [
        "Header Nav: PHIM | RẠP CGV | VÉ CỦA TÔI | CULTUREPLEX",
        "Utility Bar: Tin ưu đãi | Vé của tôi | Yêu thích | Đăng nhập",
        "CTA chính: “Mua vé ngay” → trang Phim đang chiếu",
        "Tìm kiếm toàn site: GET → phim-dang-chieu.html?q=...",
    ])
    add_heading_custom(doc, "5.2. Thành phần tương tác", 2)
    add_bullets(doc, [
        "Carousel & Rail: banner tự chạy 4.5 giây, rail phim cuộn ngang.",
        "Yêu thích ♥: lưu phim qua localStorage, badge đếm trên header.",
        "Đặt vé mock: chọn suất, ghế lưới, tính tiền, xác nhận.",
        "Tiện ích: toast thông báo, mobile menu, back-to-top, chia sẻ link.",
    ])

    # ========== 6 ==========
    add_heading_custom(doc, "6. Màu sắc và hình thức thể hiện", 1)
    add_heading_custom(doc, "6.1. Palette màu CGV", 2)
    add_table(doc,
              ["Màu", "Mã Hex", "Ứng dụng"],
              [
                  ["CGV Red", "#E71A0F", "Thương hiệu, CTA, nav active"],
                  ["Red Dark", "#B4140C", "Hover, accent"],
                  ["Cream", "#FDFCF0", "Nền nội dung"],
                  ["Dark", "#222222", "Chữ chính"],
                  ["Muted", "#666666", "Chữ phụ"],
              ],
              col_widths=[3.5, 3.5, 8.5])

    add_heading_custom(doc, "6.2. Hình thức thể hiện", 2)
    add_bullets(doc, [
        "Typography: Verdana/Arial, sans-serif; nội dung max-width 980px, căn giữa.",
        "Poster dọc 190×260px; hero slider full-width; carousel ngang.",
        "Badge độ tuổi (P/K/T13/T16/T18) và định dạng (4DX, IMAX, Gold Class).",
        "Form tương tác: tìm kiếm, lọc dropdown, chọn ghế dạng lưới.",
        "Responsive mobile + tablet; hover effect, toast, smooth scroll.",
    ])

    page_break(doc)

    # ========== 7 ==========
    add_heading_custom(doc, "7. Sitemap và User Flow", 1)
    add_heading_custom(doc, "7.1. Sitemap (6 nhánh chính)", 2)
    add_para(doc,
             "Cấu trúc điều hướng bám sát website gốc cgv.vn — phân cấp rõ ràng, dễ mở rộng:",
             first_indent=True)
    add_table(doc,
              ["Nhánh", "Nội dung con"],
              [
                  ["Trang Chủ", "Banner carousel; Phim đang chiếu; Phim sắp chiếu; Ưu đãi nổi bật"],
                  ["Phim", "Phim đang chiếu → Tìm & lọc / Chi tiết phim; Phim sắp chiếu"],
                  ["Rạp CGV", "Tất cả rạp; Tìm theo tỉnh/thành; Chi tiết rạp"],
                  ["Vé của tôi", "Vé đã mua; Phim yêu thích"],
                  ["Ưu đãi", "Tin mới & ưu đãi; Chi tiết khuyến mãi"],
                  ["Cultureplex", "Sự kiện & hoạt động; Quầy Online; eGift; Rules…"],
              ],
              col_widths=[3.5, 12])

    add_heading_custom(doc, "7.2. User Flow — Luồng chính (7 bước)", 2)
    add_para(doc,
             "Trang Chủ → Tìm/Lọc phim → Chi tiết phim → Đặt vé (chọn suất) → "
             "Chọn ghế → Xác nhận → Vé của tôi.",
             first_indent=True)
    add_heading_custom(doc, "7.3. Luồng phụ", 2)
    add_bullets(doc, [
        "Yêu thích: bấm ♥ trên trang phim → lưu vào tab Yêu thích trong “Vé của tôi”.",
        "Tìm rạp: theo tên/tỉnh → Chi tiết rạp → xem lịch chiếu.",
        "Ưu đãi: Tin ưu đãi trên header → Chi tiết khuyến mãi → áp dụng khi đặt vé.",
    ])

    # ========== 8 ==========
    add_heading_custom(doc, "8. Wireframe và thiết kế các trang chính", 1)
    add_para(doc, "Sáu trang chính được thiết kế trên Figma và hiện thực trên Front-End:", first_indent=True)
    add_table(doc,
              ["#", "Trang", "Nội dung chính", "File"],
              [
                  ["01", "Trang Chủ", "Hero, carousel, phim đang/sắp chiếu, ưu đãi nổi bật", "index.html"],
                  ["02", "Phim Đang Chiếu", "Tìm kiếm, lọc, danh sách phim dạng lưới", "phim-dang-chieu.html"],
                  ["03", "Chi Tiết Phim", "Poster, mô tả, trailer, nút đặt vé", "chi-tiet-phim.html"],
                  ["04", "Đặt Vé", "Chọn rạp, suất, ghế lưới, xác nhận", "dat-ve.html"],
                  ["05", "Tất Cả Rạp", "Tìm theo tỉnh/thành, chi tiết rạp", "tat-ca-cac-rap.html"],
                  ["06", "Vé Của Tôi", "Vé đã mua, phim yêu thích, in vé", "ve-cua-toi.html"],
              ],
              col_widths=[1.2, 3.5, 7.3, 3.5])

    # ========== 9 ==========
    add_heading_custom(doc, "9. Tính năng nổi bật (20+ tính năng tự phát triển)", 1)
    add_table(doc,
              ["#", "Tính năng", "Cách hoạt động"],
              [
                  ["1", "Tìm kiếm & lọc phim", "Theo tên/thể loại, định dạng, độ tuổi"],
                  ["2", "Sắp xếp phim", "A→Z, Z→A, khởi chiếu mới/cũ"],
                  ["3", "Đếm kết quả lọc", "Hiển thị “Hiển thị X phim”"],
                  ["4", "Tìm kiếm toàn site", "Ô tìm header → ?q= trên trang phim"],
                  ["5", "Yêu thích phim ♥", "Lưu localStorage, badge header"],
                  ["6", "Phim vừa xem", "Tự lưu khi mở chi tiết phim"],
                  ["7", "Đặt vé mock", "Chọn rạp/suất/ghế, tính tiền, lưu vé"],
                  ["8", "Vé của tôi", "Xem / xóa / in vé (mock)"],
                  ["9", "Tìm rạp theo tên", "Lọc rạp trong tỉnh đang chọn"],
                  ["10", "Chia sẻ phim", "Copy link hoặc Web Share API"],
                  ["11", "Toast thông báo", "Phản hồi yêu thích, đặt vé, copy link"],
                  ["12", "Back-to-top", "Hiện khi scroll > 320px"],
                  ["13", "Menu mobile", "Hamburger mở/đóng nav"],
                  ["14", "Banner carousel", "Tự chuyển slide 4.5s, prev/next"],
                  ["15", "Carousel phim/event", "Cuộn ngang bằng nút ‹ ›"],
                  ["16", "Badge số lượng header", "Số vé / yêu thích trên utility bar"],
                  ["17", "Responsive", "Layout tablet/mobile"],
                  ["18", "Tin mới & ưu đãi", "Listing giống cgv.vn/newsoffer"],
                  ["19", "Lọc/sắp xếp ưu đãi", "Tìm, lọc danh mục, sắp xếp ngày"],
                  ["20", "Chi tiết ưu đãi", "Trang nội dung chương trình KM"],
              ],
              col_widths=[1.2, 4.5, 9.8])

    add_heading_custom(doc, "9.1. Lưu trữ phía client (localStorage)", 2)
    add_table(doc,
              ["Key", "Dữ liệu"],
              [
                  ["cgv_fe_favorites", "Danh sách ID phim yêu thích"],
                  ["cgv_fe_tickets", "Vé đã đặt (mock)"],
                  ["cgv_fe_recent", "Tối đa 6 phim vừa xem"],
              ],
              col_widths=[5, 10.5])

    page_break(doc)

    # ========== 10 USER GUIDE ==========
    add_heading_custom(doc, "10. Hướng dẫn sử dụng (User Guide)", 1)
    add_heading_custom(doc, "10.1. Yêu cầu hệ thống", 2)
    add_table(doc,
              ["Hạng mục", "Yêu cầu"],
              [
                  ["Thiết bị", "Desktop, laptop hoặc điện thoại"],
                  ["Trình duyệt", "Chrome, Edge, Firefox hoặc Safari (bản mới)"],
                  ["JavaScript", "Bắt buộc bật"],
                  ["Mạng", "Cần mạng để tải ảnh CDN và truy cập site deploy"],
                  ["Storage", "localStorage được bật (favorites & tickets)"],
              ],
              col_widths=[4, 11.5])

    add_heading_custom(doc, "10.2. Cách truy cập website", 2)
    add_para(doc, "Cách A — Site đã publish (khuyến nghị chấm điểm):", bold=True)
    add_bullets(doc, [
        "Mở trình duyệt và vào: https://biidayy.github.io/CGV_WebsiteFE/",
        "Hoặc Vercel: https://cgv-website-fe.vercel.app",
        "Trang chủ trực tiếp: .../cgv/",
    ])
    add_para(doc, "Cách B — Xem mã nguồn GitHub:", bold=True)
    add_bullets(doc, [
        "Repo public: https://github.com/Biidayy/CGV_WebsiteFE",
        "Mã nguồn chính nằm trong thư mục cgv/",
    ])
    add_para(doc, "Cách C — Chạy local:", bold=True)
    add_para(doc, "git clone https://github.com/Biidayy/CGV_WebsiteFE.git", size=11)
    add_para(doc, "cd CGV_WebsiteFE && python -m http.server 8765", size=11)
    add_para(doc, "Mở http://127.0.0.1:8765/cgv/", size=11)

    add_heading_custom(doc, "10.3. Hướng dẫn từng bước", 2)
    steps = [
        ("10.3.1. Duyệt phim trang chủ",
         ["Mở website.", "Xem hero slider và hàng Movie Selection.",
          "Bấm poster hoặc “Xem chi tiết” để vào trang chi tiết phim."]),
        ("10.3.2. Tìm kiếm và lọc phim",
         ["Vào PHIM → Phim Đang Chiếu (hoặc Phim Sắp Chiếu).",
          "Nhập tên/thể loại vào ô tìm kiếm.",
          "Chọn lọc Thể loại / Định dạng / Độ tuổi nếu cần.",
          "Bấm Reset để xóa bộ lọc."]),
        ("10.3.3. Xem chi tiết và trailer",
         ["Mở trang chi tiết phim.",
          "Xem mô tả, thời lượng, rating, định dạng.",
          "Phát trailer (nếu có).",
          "Bấm “Mua vé ngay” để sang trang đặt vé."]),
        ("10.3.4. Đặt vé (mock)",
         ["Chọn tỉnh/thành, rạp, ngày và suất chiếu.",
          "Chọn một hoặc nhiều ghế trống trên sơ đồ ghế.",
          "Bấm “Xác nhận đặt vé” — vé lưu trong localStorage (không thanh toán thật)."]),
        ("10.3.5. Quản lý yêu thích và vé",
         ["Bấm ♥ trên poster/chi tiết để lưu yêu thích.",
          "Mở YÊU THÍCH hoặc VÉ CỦA TÔI trên thanh tiện ích.",
          "Có thể in vé hoặc xóa vé. Xóa dữ liệu trình duyệt sẽ mất các mục đã lưu."]),
        ("10.3.6. Khám phá rạp",
         ["Vào RẠP CGV → Tất Cả Các Rạp.",
          "Chọn tỉnh/thành hoặc tìm tên rạp.",
          "Xem Rạp Đặc Biệt và Rạp 3D để biết định dạng phòng chiếu."]),
    ]
    for title, items in steps:
        add_heading_custom(doc, title, 3)
        add_bullets(doc, items)

    page_break(doc)

    # ========== 11 SCREENSHOTS ==========
    add_heading_custom(doc, "11. Ảnh chụp giao diện (Screenshots)", 1)
    add_para(doc,
             "Các hình dưới đây minh họa giao diện đã hiện thực, tương ứng 5 màn hình chính "
             "trong User Guide và wireframe Figma.",
             first_indent=True)

    shots = [
        ("01-home.png", "Hình 1. Trang chủ — hero banner và hàng chọn phim"),
        ("02-now-showing.png", "Hình 2. Phim đang chiếu — danh sách kèm tìm kiếm/lọc"),
        ("03-movie-detail.png", "Hình 3. Chi tiết phim — thông tin, trailer, nút đặt vé"),
        ("04-booking.png", "Hình 4. Đặt vé mock — chọn suất và sơ đồ ghế"),
        ("05-my-tickets.png", "Hình 5. Vé của tôi / Phim yêu thích"),
    ]
    for name, caption in shots:
        add_image(doc, SHOTS / name, caption, width_inches=5.9)
        page_break(doc)

    # If last shot already page-broke, content continues — OK

    # ========== 12 ==========
    add_heading_custom(doc, "12. Giới hạn của hệ thống", 1)
    add_bullets(doc, [
        "Front-End tĩnh: không đăng nhập/đăng ký thật, không API backend.",
        "Không thanh toán thật; quy trình đặt vé chỉ mang tính mô phỏng.",
        "Vé và yêu thích chỉ lưu trong localStorage của trình duyệt.",
        "Một số mục menu (Cultureplex mở rộng, đăng nhập thành viên…) mang tính demo.",
        "Dữ liệu phim/rạp/suất là mock cố định trong js/data.js.",
        "Ảnh CDN cần kết nối Internet.",
        "Đây là sản phẩm đồ án học tập, không phải sản phẩm chính thức của CGV.",
    ])

    # ========== 13 ==========
    add_heading_custom(doc, "13. Kết luận", 1)
    add_para(doc,
             "Nhóm đã hoàn thành nghiên cứu và tái hiện giao diện CGV Cinemas Vietnam trên "
             "Front-End tĩnh, bám sát cấu trúc điều hướng website gốc cgv.vn. Hệ thống gồm "
             "các trang chính theo Figma, luồng đặt vé mock liên tục, hơn 20 tính năng tương "
             "tác phía client, giao diện responsive và bộ nhận diện thương hiệu CGV "
             "(màu đỏ #E71A0F). Website đã được deploy công khai, kèm User Guide và ảnh "
             "minh họa đầy đủ để phục vụ chấm điểm.",
             first_indent=True)
    add_para(doc,
             "Cả hai thành viên hoàn thành 100% phần việc được phân công. Nhóm xin cảm ơn "
             "thầy Nguyễn Hồng Diên và các bạn đã hỗ trợ trong quá trình thực hiện đồ án.",
             first_indent=True)

    add_heading_custom(doc, "Tài liệu tham khảo / Liên kết nộp bài", 1)
    add_table(doc,
              ["Hạng mục", "Đường dẫn"],
              [
                  ["Website (GitHub Pages)", "https://biidayy.github.io/CGV_WebsiteFE/"],
                  ["Website (Vercel)", "https://cgv-website-fe.vercel.app"],
                  ["Domain tùy chỉnh", "https://cgv-website-tkgd.is-a.dev"],
                  ["Repository", "https://github.com/Biidayy/CGV_WebsiteFE"],
                  ["Thư mục mã nguồn", "cgv/"],
                  ["Screenshot", "docs/screenshots/"],
                  ["Website gốc tham chiếu", "https://www.cgv.vn/"],
              ],
              col_widths=[5, 10.5])

    add_para(doc, "— Hết báo cáo —", size=12, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    ARTIFACT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(ARTIFACT)
    print(f"Wrote {OUT}")
    print(f"Wrote {ARTIFACT}")


if __name__ == "__main__":
    build()
